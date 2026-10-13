"""Functional container acceptance using owned fixtures and an isolated Compose project."""
import argparse
import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import time
import uuid
from urllib.parse import urlsplit

import httpx
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from reportlab.pdfgen import canvas

ROOT=Path(__file__).resolve().parents[1]


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--url',default='http://127.0.0.1:8765')
    parser.add_argument('--project',default='research-reader')
    parser.add_argument('--output',type=Path,default=ROOT/'var/docker-deployment-report.json')
    args=parser.parse_args()
    report={'fixture':'self-authored; no user documents or paid requests','checks':{},'documents':[]}
    compose=['docker','compose','--env-file','deploy/docker.env.example','-f','compose.yaml','-p',args.project]
    def run(*command):
        environment={**os.environ,'RR_HTTP_PORT':str(urlsplit(args.url).port or 8765)}
        return subprocess.check_output([*compose,*command],cwd=ROOT,env=environment,text=True,encoding='utf-8',errors='replace')
    client=httpx.Client(base_url=args.url,trust_env=False,timeout=90)
    def wait(identifier,unsupported=False):
        deadline=time.monotonic()+180
        while time.monotonic()<deadline:
            doc=client.get(f'/api/v1/documents/{identifier}').raise_for_status().json()
            if doc['parse_status']=='succeeded':return doc
            if doc['parse_status']=='failed':
                if unsupported:return doc
                raise RuntimeError(f'Fixture parse failed: {identifier}')
            time.sleep(.5)
        raise TimeoutError('Worker did not complete the fixture')
    def upload(name,data,mime,unsupported=False):
        response=client.post('/api/v1/documents',files={'file':(name,data,mime)}).raise_for_status().json()
        assert not response['duplicate']
        identifier=response['document']['id']
        report['documents'].append(identifier)
        wait(identifier,unsupported)
        assert hashlib.sha256(client.get(f'/api/v1/documents/{identifier}/original').content).digest()==hashlib.sha256(data).digest()
        return identifier
    try:
        health=client.get('/health').raise_for_status().json()
        report['checks']['api_health']=health['status']=='ok' and health['data_dir']=='/data'
        report['checks']['frontend_served']='科研文献对照阅读器' in client.get('/').text
        # Mutating requests from an unrelated web origin remain blocked.
        report['checks']['cross_origin_refused']=client.post('/api/v1/documents',headers={'Origin':'https://external.example'}).status_code==403
        output=io.BytesIO();pdf=canvas.Canvas(output,pagesize=(600,800))
        pdf.setFont('Helvetica-Bold',16);pdf.drawString(50,750,'Controlled Docker experiment '+uuid.uuid4().hex[:8])
        pdf.setFont('Helvetica-Bold',12);pdf.drawString(50,680,'1 Introduction')
        pdf.setFont('Helvetica',10);pdf.drawString(50,650,'This artificial paper evaluates an arithmetic mean of two inputs.')
        pdf.drawString(50,630,'The score is S=(x+y)/2. This text is a software test fixture.')
        pdf.setFont('Helvetica-Bold',12);pdf.drawString(50,570,'References')
        pdf.setFont('Helvetica',10);pdf.drawString(50,540,'[1] A. Example. A fictional study. Test Journal, 2025.')
        pdf.save()
        pdf_data=output.getvalue()
        identifier=upload('docker-controlled.pdf',pdf_data,'application/pdf')
        jobs=client.get('/api/v1/jobs',params={'document_id':identifier,'limit':20}).json()['jobs']
        parse=next(j for j in jobs if j['type']=='parse')
        report['grobid']=(parse.get('result') or {}).get('stats',{}).get('grobid',{})
        report['checks']['actual_worker_grobid']=report['grobid'].get('status')=='ok'
        page=client.get(f'/api/v1/documents/{identifier}/pages/0').raise_for_status().json()
        report['checks']['pdf_blocks']=bool(page['blocks'])
        report['checks']['pdf_preview']=client.get(f'/api/v1/documents/{identifier}/pages/0/preview.png').content.startswith(b'\x89PNG')
        state=f'/api/v1/documents/{identifier}/reading-state'
        client.put(state,json={'page_index':0,'view_mode':'overlay_side_by_side'}).raise_for_status()
        docx=Document();docx.add_heading('Docker DOCX sample '+uuid.uuid4().hex[:8],0);docx.add_paragraph('中文段落：用于验证容器字体和文献解析。');docx.add_paragraph('First list item',style='List Bullet')
        output=io.BytesIO();docx.save(output)
        docx_id=upload('docker-controlled.docx',output.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document',unsupported=True)
        component=json.loads(run('exec','-T','worker','python','-c',
            "from pathlib import Path; import json, pymupdf; from rr_adapters.docx_convert import DocxConverter; "
            f"source=next(Path('/data/originals/{docx_id}').iterdir()); "
            "result=DocxConverter().convert(source,Path('/tmp/component-check.pdf')); "
            "pdf=pymupdf.open('/tmp/component-check.pdf'); "
            "print(json.dumps({'chinese_text':'中文段落' in ''.join(p.get_text() for p in pdf),'pages':len(pdf)}))"))
        report['checks']['docx_component_chinese_render']=component['chinese_text'] and component['pages']>0
        image=Image.new('RGB',(1200,400),'white');draw=ImageDraw.Draw(image)
        try:font=ImageFont.truetype('arial.ttf',48)
        except OSError:font=ImageFont.truetype('DejaVuSans.ttf',48)
        draw.text((60,70),'DOCKER OCR CHECK',font=font,fill='black')
        draw.text((60,170),'CONTROLLED SAMPLE '+uuid.uuid4().hex[:6],font=font,fill='black')
        output=io.BytesIO();image.save(output,format='PNG')
        image_id=upload('docker-controlled.png',output.getvalue(),'image/png',unsupported=True)
        component=json.loads(run('exec','-T','worker','python','-c',
            "from pathlib import Path; import json, numpy as np; from PIL import Image; from rr_adapters.ocr import RapidOCRProvider; "
            "from document_core.coords import PageGeometry; "
            f"image=Image.open(next(Path('/data/originals/{image_id}').iterdir())); "
            "result=RapidOCRProvider().recognize(np.asarray(image.convert('RGB')),PageGeometry(image.width,image.height)); "
            "print(json.dumps({'text':' '.join(line.text for line in result.lines)}))"))
        report['checks']['actual_ocr_component']='DOCKER' in component['text'].upper()
        report['ocr_component']=component
        jobs=client.get('/api/v1/jobs',params={'limit':50}).json()['jobs']
        report['checks']['unwired_formats_explicitly_reported']=all(any(j['document_id']==key and j['error_code']=='unsupported_format' for j in jobs) for key in [docx_id,image_id])
        report['known_limits']=['DOCX/image adapters are callable in the image; the import worker remains PDF-only.']
        # Container replacement, not just a process restart, must retain data.
        run('up','-d','--no-build','--force-recreate','--wait','--wait-timeout','360','api','worker')
        client.close()
        client=httpx.Client(base_url=args.url,trust_env=False,timeout=90)
        report['checks']['state_persists_after_recreate']=client.get(state).json()['reading_state']['view_mode']=='overlay_side_by_side'
        report['checks']['original_persists_after_recreate']=client.get(f'/api/v1/documents/{identifier}/original').content==pdf_data
        report['runtime']=run('ps','--format','json')
        report['image_id']=run('images','--format','json')
    except Exception as error:
        report['error']=f'{type(error).__name__}: {error}'
    finally:
        report['passed']=not report.get('error') and all(report['checks'].values())
        args.output.parent.mkdir(parents=True,exist_ok=True)
        args.output.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        print(json.dumps({'passed':report['passed'],'checks':report['checks'],'error':report.get('error')},ensure_ascii=False),flush=True)
        client.close()
    return 0 if report['passed'] else 1


if __name__=='__main__':raise SystemExit(main())
